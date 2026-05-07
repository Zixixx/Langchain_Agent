from __future__ import annotations

"""文件工具：限制 Agent 只能读写项目工作区内的文件。"""

import shutil
import subprocess
import tempfile
from pathlib import Path

from langchain_core.tools import tool


TEXT_ENCODINGS = (
    "utf-8",
    "utf-8-sig",
    "gb18030",
    "gbk",
    "cp936",
    "big5",
    "shift_jis",
    "latin-1",
)

LEGACY_OFFICE_SUFFIXES = {".doc", ".ppt", ".xls"}
MODERN_OFFICE_SUFFIXES = {".docx", ".pptx", ".xlsx"}
AGENT_READ_DIRS = ("input", "output")


def _resolve_input_or_output_path(workspace: Path, user_path: str) -> Path:
    """解析 Agent 可读路径，只允许 input/ 或 output/。"""
    candidate = Path(user_path)
    if not candidate.is_absolute():
        candidate = workspace / candidate
    resolved = candidate.resolve()
    allowed_dirs = [(workspace / name).resolve() for name in AGENT_READ_DIRS]
    if not any(resolved == directory or directory in resolved.parents for directory in allowed_dirs):
        raise ValueError(f"Agent can only read files under input/ or output/: {user_path}")
    return resolved


def _resolve_output_path(workspace: Path, user_path: str) -> Path:
    """解析 Agent 可写/可执行路径，只允许 output/。"""
    candidate = Path(user_path)
    if not candidate.is_absolute():
        candidate = workspace / candidate
    target = candidate.resolve()
    output_dir = (workspace / "output").resolve()
    if target != output_dir and output_dir not in target.parents:
        raise ValueError(f"Agent can only write or execute files under output/: {user_path}")
    return target


def _read_text_with_fallback(path: Path) -> tuple[str, str]:
    """按常见编码尝试读取文本，兼容 Windows 常见 GBK/GB18030 文件。"""
    data = path.read_bytes()
    last_error: UnicodeDecodeError | None = None
    for encoding in TEXT_ENCODINGS:
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError as exc:
            last_error = exc
    if last_error:
        raise last_error
    return "", "utf-8"


def _read_pdf_text_content(path: Path) -> str:
    """用 pypdf 提取 PDF 文本内容，供文件工具和 RAG 导入共同复用。"""
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("PDF reading requires pypdf. Install it with: deps-install pypdf") from exc

    reader = PdfReader(str(path))
    parts: list[str] = []
    for page_index, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        if page_text.strip():
            parts.append(f"--- Page {page_index} ---\n{page_text.strip()}")
    text = "\n\n".join(parts).strip()
    if not text:
        raise RuntimeError(f"No extractable text found in PDF: {path}")
    return text


def _resolve_libreoffice_executable(libreoffice_path: Path | None) -> str | None:
    """优先使用 .env 指定的 LibreOffice 文件或目录，否则从 PATH 查找。"""
    if libreoffice_path:
        if libreoffice_path.is_file():
            return str(libreoffice_path)
        if libreoffice_path.is_dir():
            search_dirs = [
                libreoffice_path,
                libreoffice_path / "program",
                libreoffice_path / "Contents" / "MacOS",
            ]
            candidates = (
                "soffice.exe",
                "libreoffice.exe",
                "soffice",
                "libreoffice",
            )
            for directory in search_dirs:
                for name in candidates:
                    executable = directory / name
                    if executable.is_file():
                        return str(executable)
            raise RuntimeError(
                "LIBREOFFICE_PATH points to a directory, but no LibreOffice executable "
                f"was found in it or its program subdirectory: {libreoffice_path}"
            )
        raise RuntimeError(f"LIBREOFFICE_PATH does not exist: {libreoffice_path}")
    return shutil.which("libreoffice") or shutil.which("soffice")


def _convert_office_to_text_with_libreoffice(path: Path, libreoffice_path: Path | None) -> str:
    """用 LibreOffice headless 把老式 Office 文件转换为可读取文本。"""
    soffice = _resolve_libreoffice_executable(libreoffice_path)
    if not soffice:
        raise RuntimeError(
            "Reading .doc/.ppt/.xls requires LibreOffice. Set LIBREOFFICE_PATH in .env or run: install-libreoffice command"
        )
    with tempfile.TemporaryDirectory(prefix="agent_office_") as temp_dir:
        conversion_errors = []
        for target_format in ("txt", "csv"):
            completed = subprocess.run(
                [soffice, "--headless", "--convert-to", target_format, "--outdir", temp_dir, str(path)],
                capture_output=True,
                text=True,
                errors="replace",
                timeout=120,
                check=False,
            )
            if completed.returncode != 0:
                conversion_errors.append(completed.stderr.strip() or completed.stdout.strip())
                continue
            converted_files = sorted(Path(temp_dir).glob(f"*.{target_format}"))
            if converted_files:
                text, _encoding = _read_text_with_fallback(converted_files[0])
                return text

        error_text = "; ".join(error for error in conversion_errors if error)
        raise RuntimeError(error_text or "LibreOffice did not produce a readable text file")


def _dependency_request_message(module_name: str, package_name: str, reason: str) -> str:
    """返回给 Agent 的缺包提示，用于现代 Office 读取依赖缺失时申请安装。"""
    return (
        "DEPENDENCY_INSTALL_REQUEST\n"
        f"missing_module: {module_name}\n"
        f"suggested_package: {package_name}\n"
        f"reason: {reason}\n"
        "Call request_dependency_install with module_name and package_name, then ask the user to approve it."
    )


def _read_legacy_office_text_content(path: Path, libreoffice_path: Path | None) -> str:
    """读取老式 Office 文件；项目只支持通过 LibreOffice 转换。"""
    try:
        return _convert_office_to_text_with_libreoffice(path, libreoffice_path)
    except RuntimeError as exc:
        raise RuntimeError(f"Original error: {exc}") from exc


def _read_docx_text(path: Path) -> str:
    """读取 .docx 段落和表格文本。"""
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(_dependency_request_message("docx", "python-docx", ".docx reading requires python-docx.")) from exc

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table_index, table in enumerate(document.tables, start=1):
        parts.append(f"--- Table {table_index} ---")
        for row in table.rows:
            parts.append(",".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts).strip()


def _read_pptx_text(path: Path) -> str:
    """读取 .pptx 幻灯片中的文本框内容。"""
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError(_dependency_request_message("pptx", "python-pptx", ".pptx reading requires python-pptx.")) from exc

    presentation = Presentation(str(path))
    parts: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        slide_parts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_parts.append(shape.text.strip())
        if slide_parts:
            parts.append(f"--- Slide {slide_index} ---")
            parts.extend(slide_parts)
    return "\n".join(parts).strip()


def _read_xlsx_text(path: Path) -> str:
    """读取 .xlsx 工作表数据。"""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(_dependency_request_message("openpyxl", "openpyxl", ".xlsx reading requires openpyxl.")) from exc

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[str] = []
    for worksheet in workbook.worksheets:
        blocks.append(f"--- Sheet: {worksheet.title} | rows={worksheet.max_row} | columns={worksheet.max_column} ---")
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            blocks.append(",".join("" if value is None else str(value) for value in row))
    workbook.close()
    return "\n".join(blocks).strip()


def _read_modern_office_text(path: Path) -> str:
    """按扩展名读取现代 Office Open XML 文件。"""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _read_docx_text(path)
    if suffix == ".pptx":
        return _read_pptx_text(path)
    if suffix == ".xlsx":
        return _read_xlsx_text(path)
    raise RuntimeError(f"Unsupported modern Office file type: {path.name}")


def build_file_tools(workspace: Path, libreoffice_path: Path | None = None):
    """构造目录列表、文本读取、PDF/老式 Office 读取和文本写入工具。"""
    workspace = workspace.resolve()

    @tool
    def list_directory(path: str = ".") -> str:
        """List input/output directories. Use "." to show only the Agent-visible roots."""
        if path in {"", "."}:
            target = workspace
        else:
            try:
                target = _resolve_input_or_output_path(workspace, path)
            except ValueError as exc:
                return str(exc)
        if not target.exists():
            return f"Path does not exist: {path}"
        if not target.is_dir():
            return f"Not a directory: {path}"
        if target == workspace:
            entries = []
            for name in AGENT_READ_DIRS:
                item = workspace / name
                if item.exists():
                    entries.append(f"dir\t{name}")
            return "\n".join(entries) or "(empty)"
        entries = []
        for item in sorted(target.iterdir(), key=lambda p: (not p.is_dir(), p.name.lower())):
            kind = "dir" if item.is_dir() else "file"
            entries.append(f"{kind}\t{item.relative_to(workspace)}")
        return "\n".join(entries) or "(empty)"

    @tool
    def read_text_file(path: str) -> str:
        """Read a text file under input/ or output/. Tries UTF-8, GB18030, GBK, Big5 and other common encodings."""
        try:
            target = _resolve_input_or_output_path(workspace, path)
        except ValueError as exc:
            return str(exc)
        if not target.exists():
            return f"File does not exist: {path}"
        if not target.is_file():
            return f"Not a file: {path}"
        suffix = target.suffix.lower()
        if suffix == ".pdf":
            return "PDF files are binary. Use read_pdf_text for .pdf files."
        if suffix in LEGACY_OFFICE_SUFFIXES:
            return "Legacy Office files are binary. Use read_legacy_office_text for .doc/.ppt/.xls files."
        if suffix in MODERN_OFFICE_SUFFIXES:
            return "Modern Office files are zipped XML packages. Use read_modern_office_text for .docx/.pptx/.xlsx files."
        try:
            text, encoding = _read_text_with_fallback(target)
        except UnicodeDecodeError as exc:
            return f"File could not be decoded with common text encodings: {path}. Error: {exc}"
        return f"[encoding={encoding}]\n" + text

    @tool
    def read_pdf_text(path: str) -> str:
        """Extract text from a PDF file under input/ or output/. Use this for .pdf papers and reports."""
        try:
            target = _resolve_input_or_output_path(workspace, path)
        except ValueError as exc:
            return str(exc)
        if not target.exists():
            return f"File does not exist: {path}"
        if not target.is_file():
            return f"Not a file: {path}"
        if target.suffix.lower() != ".pdf":
            return f"Not a PDF file: {path}"
        try:
            text = _read_pdf_text_content(target)
        except Exception as exc:
            return f"Failed to read PDF: {path}. Error: {exc}"

        return text

    @tool
    def read_legacy_office_text(path: str) -> str:
        """Extract text from legacy Office files under input/ or output/: .doc, .ppt and .xls."""
        try:
            target = _resolve_input_or_output_path(workspace, path)
        except ValueError as exc:
            return str(exc)
        if not target.exists():
            return f"File does not exist: {path}"
        if not target.is_file():
            return f"Not a file: {path}"
        suffix = target.suffix.lower()
        try:
            if suffix in LEGACY_OFFICE_SUFFIXES:
                text = _read_legacy_office_text_content(target, libreoffice_path)
            else:
                return f"Unsupported legacy Office file type: {path}. Supported: .doc, .ppt, .xls"
        except Exception as exc:
            return f"Failed to read legacy Office file: {path}. Error: {exc}"
        if not text.strip():
            return f"No extractable text found in file: {path}"
        return text.strip()

    @tool
    def read_modern_office_text(path: str) -> str:
        """Extract text from modern Office files under input/ or output/: .docx, .pptx and .xlsx."""
        try:
            target = _resolve_input_or_output_path(workspace, path)
        except ValueError as exc:
            return str(exc)
        if not target.exists():
            return f"File does not exist: {path}"
        if not target.is_file():
            return f"Not a file: {path}"
        suffix = target.suffix.lower()
        try:
            if suffix in MODERN_OFFICE_SUFFIXES:
                text = _read_modern_office_text(target)
            else:
                return f"Unsupported modern Office file type: {path}. Supported: .docx, .pptx, .xlsx"
        except Exception as exc:
            return f"Failed to read modern Office file: {path}. Error: {exc}"
        if not text.strip():
            return f"No extractable text found in file: {path}"
        return text.strip()

    @tool
    def copy_file(source_path: str = "", output_path: str = "") -> str:
        """Copy an existing input/output file to a new file under output/. Refuses to overwrite."""
        if not source_path or not output_path:
            return "Usage: copy_file(source_path, output_path). Both arguments are required."
        try:
            source = _resolve_input_or_output_path(workspace, source_path)
        except ValueError as exc:
            return str(exc)
        if not source.exists():
            return f"Source file does not exist: {source_path}"
        if not source.is_file():
            return f"Source is not a file: {source_path}"

        target = _resolve_output_path(workspace, output_path)
        if target.exists():
            return f"Refused to overwrite existing file: {target.relative_to(workspace)}. Use a new filename."

        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)
        return (
            f"Copied {source.relative_to(workspace)} to {target.relative_to(workspace)} "
            f"({target.stat().st_size} bytes)."
        )

    @tool
    def rename_file(path: str = "", new_path: str = "") -> str:
        """Rename or move one existing output/ file to another new output/ path. Refuses to overwrite."""
        if not path or not new_path:
            return "Usage: rename_file(path, new_path). Both arguments are required."
        try:
            source = _resolve_output_path(workspace, path)
            target = _resolve_output_path(workspace, new_path)
        except ValueError as exc:
            return str(exc)
        if not source.is_file():
            return f"Output file does not exist: {path}"
        if target.exists():
            return f"Refused to overwrite existing file: {target.relative_to(workspace)}. Use a new filename."
        target.parent.mkdir(parents=True, exist_ok=True)
        source.rename(target)
        return f"Renamed {source.relative_to(workspace)} to {target.relative_to(workspace)}"

    @tool
    def delete_file(path: str = "") -> str:
        """Delete one existing file under output/. Directories and input/ files are refused."""
        if not path:
            return "Usage: delete_file(path). Missing required path."
        try:
            target = _resolve_output_path(workspace, path)
        except ValueError as exc:
            return str(exc)
        if not target.is_file():
            return f"Output file does not exist: {path}"
        target.unlink()
        return f"Deleted {target.relative_to(workspace)}"

    @tool
    def create_directory(path: str = "") -> str:
        """Create a new directory under output/. Refuses to overwrite existing paths."""
        if not path:
            return "Usage: create_directory(path). Missing required path."
        try:
            target = _resolve_output_path(workspace, path)
        except ValueError as exc:
            return str(exc)
        if target.exists():
            return f"Refused to overwrite existing path: {target.relative_to(workspace)}. Use a new directory."
        target.mkdir(parents=True, exist_ok=True)
        return f"Created directory: {target.relative_to(workspace)}"

    @tool
    def copy_directory(source_path: str = "", output_path: str = "") -> str:
        """Copy an input/output directory to a new directory under output/. Refuses to overwrite."""
        if not source_path or not output_path:
            return "Usage: copy_directory(source_path, output_path). Both arguments are required."
        try:
            source = _resolve_input_or_output_path(workspace, source_path)
            target = _resolve_output_path(workspace, output_path)
        except ValueError as exc:
            return str(exc)
        if not source.is_dir():
            return f"Source directory does not exist: {source_path}"
        if target.exists():
            return f"Refused to overwrite existing path: {target.relative_to(workspace)}. Use a new directory."
        shutil.copytree(source, target)
        file_count = sum(1 for item in target.rglob("*") if item.is_file())
        return f"Copied directory {source.relative_to(workspace)} to {target.relative_to(workspace)} ({file_count} file(s))."

    @tool
    def rename_directory(path: str = "", new_path: str = "") -> str:
        """Rename or move one output/ directory to another new output/ path. Refuses to overwrite."""
        if not path or not new_path:
            return "Usage: rename_directory(path, new_path). Both arguments are required."
        try:
            source = _resolve_output_path(workspace, path)
            target = _resolve_output_path(workspace, new_path)
        except ValueError as exc:
            return str(exc)
        if not source.is_dir():
            return f"Output directory does not exist: {path}"
        if target.exists():
            return f"Refused to overwrite existing path: {target.relative_to(workspace)}. Use a new directory."
        if source == target or source in target.parents:
            return "Refused to move a directory into itself or its child directory."
        target.parent.mkdir(parents=True, exist_ok=True)
        source.rename(target)
        return f"Renamed directory {source.relative_to(workspace)} to {target.relative_to(workspace)}"

    @tool
    def delete_directory(path: str = "") -> str:
        """Delete one directory under output/. Refuses to delete output/ itself."""
        if not path:
            return "Usage: delete_directory(path). Missing required path."
        try:
            target = _resolve_output_path(workspace, path)
        except ValueError as exc:
            return str(exc)
        output_dir = (workspace / "output").resolve()
        if target == output_dir:
            return "Refused to delete the output/ root directory."
        if not target.is_dir():
            return f"Output directory does not exist: {path}"
        file_count = sum(1 for item in target.rglob("*") if item.is_file())
        shutil.rmtree(target)
        return f"Deleted directory {path} ({file_count} file(s))."

    @tool
    def write_text_file(path: str = "", content: str = "") -> str:
        """Write a new UTF-8 text file to the output directory. Refuses to overwrite existing files."""
        if not path:
            return "Usage: write_text_file(path, content). Missing required path, for example: output/script.py"
        target = _resolve_output_path(workspace, path)
        if target.exists():
            return f"Refused to overwrite existing file: {target.relative_to(workspace)}. Use a new filename."
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(content, encoding="utf-8")
        return f"Wrote {len(content)} chars to {target.relative_to(workspace)}"

    @tool
    def append_text_file(path: str = "", content: str = "") -> str:
        """Append UTF-8 text to a generated file under output/. Use this to build long scripts in chunks."""
        if not path:
            return "Usage: append_text_file(path, content). Missing required path, for example: output/script.py"
        target = _resolve_output_path(workspace, path)
        target.parent.mkdir(parents=True, exist_ok=True)
        with target.open("a", encoding="utf-8", newline="") as file:
            file.write(content)
        return f"Appended {len(content)} chars to {target.relative_to(workspace)}"

    @tool
    def replace_text_file(path: str = "", old_text: str = "", new_text: str = "") -> str:
        """Replace one exact text fragment in a generated file under output/. Fails if the match is missing or ambiguous."""
        if not path:
            return "Usage: replace_text_file(path, old_text, new_text). Missing required path."
        target = _resolve_output_path(workspace, path)
        if not target.is_file():
            return f"Output file does not exist: {path}"
        if not old_text:
            return "Refused to replace empty old_text."

        content = target.read_text(encoding="utf-8")
        match_count = content.count(old_text)
        if match_count == 0:
            return "old_text was not found. Provide the exact current text fragment to replace."
        if match_count > 1:
            return f"old_text matched {match_count} times. Provide a more specific fragment and call replace_text_file once."

        updated = content.replace(old_text, new_text, 1)
        target.write_text(updated, encoding="utf-8")
        return (
            f"Replaced one fragment in {target.relative_to(workspace)} "
            f"({len(old_text)} chars -> {len(new_text)} chars)."
        )

    return [
        list_directory,
        read_text_file,
        read_pdf_text,
        read_legacy_office_text,
        read_modern_office_text,
        copy_file,
        rename_file,
        delete_file,
        create_directory,
        copy_directory,
        rename_directory,
        delete_directory,
        write_text_file,
        append_text_file,
        replace_text_file,
    ]
